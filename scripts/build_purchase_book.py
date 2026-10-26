"""Turn Purchase 2026-27.xlsx into dashboard JSON.

Bank account, IFSC and bank name on the Payments sheet are never written out.
Rerun: python3 scripts/build_purchase_book.py
"""
from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from pathlib import Path

SRC = Path("/Users/anshulraibole/Downloads/Purchase 2026-27.xlsx")
OUT_DIR = Path(__file__).resolve().parents[1] / "src" / "lib"

TYPOS = (
    ("indudtries", "industries"),
    ("industrirs", "industries"),
    ("indusries", "industries"),
    ("grassim", "grasim"),
    ("industris", "industries"),
)


def fingerprint(name: str) -> str:
    s = name.lower().replace("&", " and ")
    s = re.sub(r"[^a-z0-9 ]+", " ", s)
    for src, dst in TYPOS:
        s = re.sub(rf"\b{src}\b", dst, s)
    s = re.sub(r"\b(private|pvt|limited|ltd)\b", " ", s)
    s = re.sub(r"\bind\b", "industries", s)
    return re.sub(r"\s+", " ", s).strip()


def tidy(name: str) -> str:
    return re.sub(r"\s+", " ", name).strip().rstrip(".")


def text(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def num(value):
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    raw = str(value).replace(",", "").replace("/-", "").strip()
    try:
        return round(float(raw), 2)
    except ValueError:
        return None


def colour(cell) -> str:
    try:
        rgb = str(cell.fill.fgColor.rgb)
    except Exception:
        return ""
    rgb = rgb.upper()
    if rgb in ("00000000", "NONE", "0"):
        return ""
    return rgb[-6:]


def pay_status(rgb: str) -> str:
    return {
        "00FFFF": "updated",
        "FF6F31": "input",
        "FF9900": "input",
        "00FF00": "done",
        "FF0000": "stopped",
        "FFFF00": "pending",
    }.get(rgb, "open")


def req_status(rgb: str) -> str:
    return {
        "00FF00": "received",
        "00FF0D": "received",
        "FF9900": "dispatched",
        "FFFF00": "approved",
        "FF00FF": "proforma",
        "FF0000": "stopped",
    }.get(rgb, "open")


def is_chlorine(item: str) -> bool:
    s = item.lower()
    return "liquid chlorine" in s or "liquide chlorine" in s


def parse_date(value):
    if value is None:
        return None
    if hasattr(value, "year") and hasattr(value, "month"):
        return value.year, value.month, f"{value.day:02d}.{value.month:02d}.{value.year}"
    m = re.match(r"(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})", str(value).strip())
    if not m:
        return None
    d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
    if y < 100:
        y += 2000
    return y, mo, f"{d:02d}.{mo:02d}.{y}"


def in_sep_2026(value) -> bool:
    parsed = parse_date(value)
    return bool(parsed and parsed[0] == 2026 and parsed[1] == 9)


def in_fy_2026(value) -> bool:
    parsed = parse_date(value)
    if not parsed:
        return False
    y, mo, _ = parsed
    return (y == 2026 and mo >= 4) or (y == 2027 and mo <= 3)


def show_date(value) -> str:
    parsed = parse_date(value)
    if parsed:
        return parsed[2]
    return text(value)


def build_aliases(names: Counter) -> dict:
    groups = defaultdict(list)
    for name, count in names.items():
        groups[fingerprint(name)].append((count, name))
    # "petrochemicals" and "petrochemical" are the same party.
    for key in list(groups):
        stem = key[:-1] if key.endswith("s") else ""
        if stem and stem in groups and stem != key:
            groups[stem].extend(groups.pop(key))
            groups[stem].append((0, key))  # remember the absorbed fingerprint
    aliases = {}
    for key, variants in groups.items():
        spellings = [(c, n) for c, n in variants if c > 0]
        absorbed = [n for c, n in variants if c == 0]
        if not spellings:
            continue
        best = tidy(max(spellings, key=lambda pair: pair[0])[1])
        aliases[key] = best
        for extra in absorbed:
            aliases[extra] = best
    return aliases


def main() -> None:
    import openpyxl

    book = openpyxl.load_workbook(SRC, data_only=False)
    values = openpyxl.load_workbook(SRC, data_only=True)

    purchase = book["Purchase "]
    purchase_v = values["Purchase "]
    names = Counter()
    for row in range(3, purchase_v.max_row + 1):
        supplier = text(purchase_v.cell(row, 4).value)
        if supplier:
            names[supplier] += 1
    aliases = build_aliases(names)

    def canon(name: str) -> str:
        cleaned = tidy(name)
        if not cleaned:
            return ""
        return aliases.get(fingerprint(cleaned), cleaned)

    purchases = []
    sep_ex_chlorine = 0.0
    for row in range(3, purchase_v.max_row + 1):
        supplier = text(purchase_v.cell(row, 4).value)
        item = text(purchase_v.cell(row, 5).value)
        amount = num(purchase_v.cell(row, 8).value)
        if not supplier and not item and amount is None:
            continue
        chlorine = is_chlorine(item)
        if amount is not None and in_sep_2026(purchase_v.cell(row, 2).value) and not chlorine:
            sep_ex_chlorine += amount
        purchases.append({
            "company": text(purchase_v.cell(row, 1).value),
            "date": show_date(purchase_v.cell(row, 2).value),
            "invoice": text(purchase_v.cell(row, 3).value),
            "supplier": canon(supplier),
            "item": item,
            "place": text(purchase_v.cell(row, 6).value),
            "qty": text(purchase_v.cell(row, 7).value),
            "amount": amount,
            "status": pay_status(colour(purchase.cell(row, 8))),
            "vehicle": text(purchase_v.cell(row, 10).value),
            "received": show_date(purchase_v.cell(row, 13).value),
            "chlorine": chlorine,
        })

    payments = []
    pay = book["Payments"]
    pay_v = values["Payments"]
    counts = Counter()
    for row in range(3, pay_v.max_row + 1):
        amount = num(pay_v.cell(row, 5).value)
        beneficiary = text(pay_v.cell(row, 3).value)
        if amount is None and not beneficiary:
            continue
        status = pay_status(colour(pay.cell(row, 5)))
        counts[status] += 1
        payments.append({
            "company": text(pay_v.cell(row, 1).value),
            "date": show_date(pay_v.cell(row, 2).value),
            "beneficiary": beneficiary,
            "amount": amount,
            "status": status,
            "inputer": text(pay_v.cell(row, 9).value),
            "remarks": text(pay_v.cell(row, 10).value),
        })

    paraffin = []
    para = values["Paraffin Payment Details"]
    for row in range(2, para.max_row + 1):
        party = text(para.cell(row, 4).value)
        if not party:
            continue
        paraffin.append({
            "company": text(para.cell(row, 1).value),
            "date": show_date(para.cell(row, 2).value),
            "invoice": text(para.cell(row, 3).value),
            "party": canon(party),
            "qty": num(para.cell(row, 5).value),
            "price": num(para.cell(row, 6).value),
            "taxable": num(para.cell(row, 7).value),
            "invoiceValue": num(para.cell(row, 8).value),
            "tds": num(para.cell(row, 9).value),
            "payable": num(para.cell(row, 12).value),
            "due": show_date(para.cell(row, 13).value),
            "cheque": text(para.cell(row, 14).value),
            "remarks": text(para.cell(row, 15).value),
        })

    requirements = []
    req = book["Requirements"]
    req_v = values["Requirements"]
    company = text(req_v.cell(2, 1).value) or "MADAN"
    for row in range(3, req_v.max_row + 1):
        named = text(req_v.cell(row, 1).value)
        if named and len(named) < 40:
            company = named
        item = text(req_v.cell(row, 4).value)
        if not item or item.upper() in {"QUANTITY", "ITEM"}:
            continue
        requirements.append({
            "company": company,
            "date": show_date(req_v.cell(row, 2).value),
            "by": text(req_v.cell(row, 3).value),
            "item": item,
            "qty": text(req_v.cell(row, 5).value),
            "supplier": canon(text(req_v.cell(row, 7).value)),
            "invoice": text(req_v.cell(row, 9).value),
            "amount": text(req_v.cell(row, 10).value),
            "vehicle": text(req_v.cell(row, 11).value),
            "dispatched": show_date(req_v.cell(row, 13).value),
            "received": show_date(req_v.cell(row, 15).value),
            "status": req_status(colour(req.cell(row, 4))),
        })

    dispatch_rows = []
    disp = values["Dispatch "]
    for row in range(2, disp.max_row + 1):
        document = text(disp.cell(row, 3).value)
        destination = text(disp.cell(row, 1).value)
        if not document and not destination:
            continue
        dispatch_rows.append({
            "destination": destination,
            "date": show_date(disp.cell(row, 2).value),
            "document": document,
            "sender": text(disp.cell(row, 4).value),
            "from": text(disp.cell(row, 5).value),
            "supplier": canon(text(disp.cell(row, 6).value)),
            "vehicle": text(disp.cell(row, 8).value),
            "receiver": text(disp.cell(row, 9).value),
            "received": show_date(disp.cell(row, 10).value),
            "remarks": text(disp.cell(row, 11).value),
        })

    def site_rows(sheet: str, cols: dict) -> list:
        ws = values[sheet]
        out = []
        for row in range(2, ws.max_row + 1):
            when = show_date(ws.cell(row, cols["date"]).value)
            if not when:
                continue
            out.append({key: (show_date(ws.cell(row, idx).value) if key == "date" else (
                num(ws.cell(row, idx).value) if key in {"qty", "rate", "amount", "price"} else text(ws.cell(row, idx).value)
            )) for key, idx in cols.items()})
        return out

    kolkata = []
    kws = values["Kolkata"]
    for row in range(2, kws.max_row + 1):
        when = show_date(kws.cell(row, 1).value)
        if not when:
            continue
        kolkata.append({
            "date": when,
            "invoice": text(kws.cell(row, 2).value),
            "supplier": canon(text(kws.cell(row, 3).value)),
            "material": text(kws.cell(row, 4).value),
            "qty": num(kws.cell(row, 5).value),
            "rate": num(kws.cell(row, 6).value),
            "amount": num(kws.cell(row, 7).value),
        })

    shd = []
    sws = values["SHD"]
    for row in range(2, sws.max_row + 1):
        when = show_date(sws.cell(row, 1).value)
        if not when:
            continue
        shd.append({
            "date": when,
            "company": text(sws.cell(row, 2).value),
            "gravity": text(sws.cell(row, 3).value),
            "qty": text(sws.cell(row, 4).value),
            "price": text(sws.cell(row, 5).value),
            "amount": num(sws.cell(row, 6).value),
            "truck": text(sws.cell(row, 7).value),
        })

    def month_sum(rows, pred) -> tuple[float, int]:
        total = 0.0
        count = 0
        for row in rows:
            if pred(row["date"]) and isinstance(row.get("amount"), (int, float)):
                total += row["amount"]
                count += 1
        return round(total, 2), count

    # Last numeric marine running balance on column O.
    marine_balance = None
    marine_as_of = ""
    marine_company = ""
    pws = values["Purchase "]
    for row in range(2, pws.max_row + 1):
        value = pws.cell(row, 15).value
        if isinstance(value, (int, float)):
            marine_balance = round(float(value), 2)
            marine_as_of = show_date(pws.cell(row, 2).value)
            marine_company = text(pws.cell(row, 1).value)

    shd_sep, shd_sep_n = month_sum(shd, in_sep_2026)
    kol_sep, kol_sep_n = month_sum(kolkata, in_sep_2026)
    shd_fy, _ = month_sum(shd, in_fy_2026)
    kol_fy, _ = month_sum(kolkata, in_fy_2026)

    meta = {
        "marineBalance": marine_balance,
        "marineAsOf": marine_as_of,
        "marineCompany": marine_company,
        "shdSeptember": shd_sep,
        "shdSeptemberRows": shd_sep_n,
        "kolkataSeptember": kol_sep,
        "kolkataSeptemberRows": kol_sep_n,
        "shdFy": shd_fy,
        "kolkataFy": kol_fy,
        "purchaseSeptemberExChlorine": round(sep_ex_chlorine, 2),
        "paymentCounts": dict(counts),
    }
    payload = {
        "purchases": purchases,
        "payments": payments,
        "paraffin": paraffin,
        "requirements": requirements,
        "dispatch": dispatch_rows,
        "kolkata": kolkata,
        "shd": shd,
    }
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUT_DIR / "supplierAliases.json").write_text(json.dumps(aliases, ensure_ascii=False, indent=2) + "\n")
    (OUT_DIR / "purchaseBookMeta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2) + "\n")
    (OUT_DIR / "purchaseBook.json").write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")))
    print("aliases", len(aliases))
    print("purchases", len(purchases), "sep ex chlorine", round(sep_ex_chlorine, 2))
    print("payments", len(payments), dict(counts))
    print("paraffin", len(paraffin), "requirements", len(requirements), "dispatch", len(dispatch_rows))
    print("kolkata", len(kolkata), "sep", kol_sep, "shd", len(shd), "sep", shd_sep)
    print("marine", marine_balance, marine_as_of, marine_company)


if __name__ == "__main__":
    main()
